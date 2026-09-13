#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extrait la carte des vins depuis l'Excel du restaurant vers vins.json."""

import json, re, unicodedata, sys
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Vin Monarque OG.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "vins.json"

TOP = {"Vins au verre", "Bulles", "Vin blanc", "Vin rosé", "Vin orange", "Vin rouge"}

COULEUR_VERRE = {
    "Apéritif": "apéritif",
    "Bulles": "bulles",
    "Blanc": "blanc",
    "Rouge": "rouge",
}

MILL_RE = re.compile(r"\b(1[89]\d{2}|20\d{2})\b")
PAREN_RE = re.compile(r"\(([^)]*)\)")


def norm(s):
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).strip()


def slug(s):
    s = unicodedata.normalize("NFD", (s or "").lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f == 0:
        return None
    return round(f, 2)


# --- lecture brute -----------------------------------------------------------
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.active

wbf = openpyxl.load_workbook(SRC)  # avec styles : la taille de police donne le niveau de titre
wsf = wbf.active


def font_size(i):
    try:
        f = wsf.cell(row=i, column=2).font
        return float(f.size or 0), bool(f.bold)
    except Exception:
        return 0.0, False


rows = []  # (idx, code, label, cepages, agence, prix_formule, prix_carte)
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    code = norm(str(r[0])) if r[0] is not None else ""
    label = norm(r[1] if isinstance(r[1], str) else (str(r[1]) if r[1] else ""))
    cep = norm(r[2] if isinstance(r[2], str) else "")
    agence = norm(r[3] if isinstance(r[3], str) else "")
    pf = num(r[4])
    pc = num(r[5])
    if not label and not code:
        continue
    rows.append(dict(idx=i, code=code, label=label, cepages=cep, agence=agence,
                     prix_formule=pf, prix_carte=pc))

# un en-tête de section = libellé sans prix carte et sans code SAQ
for k, r in enumerate(rows):
    r["is_header"] = bool(r["label"]) and not r["code"] and r["prix_carte"] is None
    # ligne vide / séparateur
    if r["label"] in ("", "-"):
        r["is_header"] = False

for r in rows:
    size, bold = font_size(r["idx"])
    # 24 pt = catégorie, 18 pt = région, 14 pt = sous-région
    r["niveau"] = 0 if size >= 22 else (1 if size >= 16 else 2)


# --- parsing d'un libellé de vin --------------------------------------------
FORMAT_RE = re.compile(r"\b(magnum|jéroboam|jeroboam|double magnum|demi[- ]bouteille|375\s?ml)\b", re.I)
TAIL_NOISE = re.compile(
    r"^(rosé|orange|blanc|rouge|demi[- ]sec|moelleux|sec|doux|nature|bio|magnum)\b[,\s]*", re.I)


GRAPES = set("""chardonnay chardo pinot noir blanc gris beurot meunier meslier savagnin aligote
gamay syrah grenache gren garnacha carignan cinsault mourvedre mourv counoise vaccarese terret
merlot cabernet cab sauvignon sauv franc frc malbec petit verdot cot tannat negrette duras fer
sangiovese sangio nebbiolo barbera dolcetto vespolina freisa grignolino ruche croatina corvina
rondinella molinara lagrein teroldego schiava trebbiano pecorino falanghina fiano greco grillo
catarratto carricante nerello mascalese capuccio nero avola montepulciano aglianico primitivo
negroamaro cannonau vermentino verdicchio garganega glera prosecco lambrusco grasparossa
tempranillo garnacho graciano mazuelo mencia caino espadeiro listan negro vijariego manto
mantonegro callet monastrell bobal xarel lo xarello macabeo maccabeu parellada godello albarino
verdejo treixadura loureiro palomino pedro ximenez moscatel baga bical cercial encruzado arinto
castelao aragonez moreto alicante bouschet alfrocheiro touriga nacional franca tinto roriz cao
francisca sousao barroca riesling gruner veltliner gruner sylvaner gewurztraminer gewurz muscat
auxerrois spatburgunder blaufrankisch blaufrnkisch zweigelt st laurent furmint harslevelu
chenin melon romorantin folle grolleau pineau aunis negrette sacy saint bris
marsanne mars roussanne rouss viognier clairette bourboulenc picardan picpoul maccabeo rolle
ugni colombard semillon muscadelle xinomavro agiorgitiko assyrtiko robola moschofilero malagousia
niellucciu sciacarellu sciacarello barbarossa jacquere altesse mondeuse persan poulsard pouls
trousseau ploussard tamjanika rkatsiteli saperavi frontenac marquette vidal seyval petite perle
zinfandel carmenere bonarda torrontes criolla pais cinsaut tinta francisco tintilia""".split())

NOISE_TAIL = set("rose rosé foncé fonce demi sec moelleux doux tranquille".split())
AMBIG = {"blanc", "blancs", "blc", "noir", "noirs", "gris"}

PROTECT = set("""champagne brut cru grand premier 1er nature extra vino vin
vins france italie espagne portugal reserve cuvee bourgogne rose dosage zero dozaggio spumante
crmant cremant loire bugey cerdon bairrada franciacorta touraine sec demi rosato solera anfora
bulles physis perpetum eclipsia renaissance mepetit prisme avizoise cramant bouzy""".split())


def strip_grape_tail(before, cep_col):
    """Quand le millésime est absent, retire les cépages collés en fin de libellé."""
    ctoks = set(slug(cep_col).split())
    toks = before.split()

    def is_grape(t):
        s = slug(t).strip()
        if not s or s in PROTECT:
            return False
        if s in GRAPES:
            return True
        for c in ctoks:
            if len(c) >= 4 and (s.startswith(c) or c.startswith(s)) and len(s) >= 4:
                return True
        return False

    removed = []
    while len(toks) > 2:
        t = toks[-1].strip(",&")
        s = slug(t)
        prev = slug(toks[-2]) if len(toks) >= 2 else ""
        parts_s = s.split()
        if s in ("et", "&", "") or (parts_s and all(p in NOISE_TAIL for p in parts_s)):
            removed.insert(0, toks.pop())
            continue
        if s in AMBIG:
            if prev in ("de", "des", "du"):   # « Blanc de Blancs » : on garde
                break
            removed.insert(0, toks.pop())
            continue
        if is_grape(t):
            removed.insert(0, toks.pop())
        else:
            break
    return norm(" ".join(toks)).strip(" ,"), norm(" ".join(removed)).strip(" ,")


def parse_label(label, couleur_hint="", cep_col=""):
    s = norm(label)
    parens = [norm(p) for p in PAREN_RE.findall(s)]
    s_clean = norm(PAREN_RE.sub(" ", s))

    fmt = ""
    m = FORMAT_RE.search(s_clean)
    if m:
        fmt = m.group(1).capitalize()
        s_clean = norm(s_clean.replace(m.group(0), " "))

    millesime = ""
    mm = list(MILL_RE.finditer(s_clean))
    if mm:
        m = mm[-1]
        millesime = m.group(1)
        before, after = s_clean[: m.start()], s_clean[m.end():]
    else:
        m2 = re.search(r"\b(NM|NV)\b", s_clean)
        if m2:
            millesime = "NM"
            before, after = s_clean[: m2.start()], s_clean[m2.end():]
        else:
            before, after = s_clean, ""

    before = norm(before).strip(" ,")
    after = norm(after).strip(" ,")

    if not millesime:
        before, grape_tail = strip_grape_tail(before, cep_col)
        after = norm((grape_tail + " " + after)).strip(" ,")

    # cépages en queue de libellé
    tail = after
    prev = None
    while tail != prev:
        prev = tail
        tail = norm(TAIL_NOISE.sub("", tail)).strip(" ,")

    parts = [p.strip() for p in before.split(",") if p.strip()]
    producteur = parts[0] if parts else before
    appellation = parts[-1] if len(parts) > 1 else ""
    cuvee = ", ".join(parts[1:-1]) if len(parts) > 2 else ""

    # pays / région entre parenthèses (on ignore les mentions de dégorgement)
    lieux = [p for p in parens if not re.match(r"^\s*dég", p, re.I)]
    origine = " · ".join(lieux)

    return dict(producteur=producteur, cuvee=cuvee, appellation=appellation,
                millesime=millesime, origine=origine, format=fmt,
                cepages_label=tail)


# --- parcours ----------------------------------------------------------------
vins = []
stack = ["", "", ""]  # [catégorie, région, sous-région]

for r in rows:
    name = r["label"]
    if r["is_header"]:
        n = r["niveau"]
        stack[n] = name
        for j in range(n + 1, 3):
            stack[j] = ""
        continue

    if r["prix_carte"] is None and r["prix_formule"] is None:
        continue  # séparateur
    if not r["label"]:
        continue

    chemin = [x for x in stack if x]
    top = stack[0]
    feuille = chemin[-1] if chemin else ""

    au_verre = top == "Vins au verre"

    if au_verre:
        couleur = COULEUR_VERRE.get(feuille, "")
        if feuille == "Rosé et orange":
            low = slug(r["label"])
            couleur = "orange" if " orange" in " " + low else "rosé"
        region = ""
    else:
        couleur = {"Bulles": "bulles", "Vin blanc": "blanc", "Vin rosé": "rosé",
                   "Vin orange": "orange", "Vin rouge": "rouge"}.get(top, "")
        region = " · ".join(chemin[1:])

    p = parse_label(r["label"], couleur, r["cepages"])

    cep_src = r["cepages"]
    # la colonne C mélange cépages et notes techniques : on garde tout, c'est utile en service
    notes = cep_src

    v = dict(
        id=None,
        code=r["code"] or "",
        label=r["label"],
        producteur=p["producteur"],
        cuvee=p["cuvee"],
        appellation=p["appellation"],
        millesime=p["millesime"],
        origine=p["origine"],
        region=region,
        couleur=couleur,
        format=p["format"],
        cepages=p["cepages_label"] or cep_src,
        notes=notes,
        agence=r["agence"],
        prix_bouteille=None if au_verre else r["prix_carte"],
        prix_verre=r["prix_carte"] if au_verre else None,
        section=" · ".join(chemin),
        ligne=r["idx"],
    )
    vins.append(v)

# --- fusion des vins au verre avec leur fiche bouteille ----------------------
def key(v):
    base = slug(v["producteur"]) + "|" + slug(v["cuvee"] or v["appellation"])
    return base

bouteilles = [v for v in vins if v["prix_bouteille"] is not None]
verres = [v for v in vins if v["prix_verre"] is not None]
index = {}
for v in bouteilles:
    index.setdefault(key(v), []).append(v)

final = list(bouteilles)
for v in verres:
    cands = index.get(key(v), [])
    match = None
    for c in cands:
        if not v["millesime"] or not c["millesime"] or v["millesime"] == c["millesime"]:
            match = c
            break
    if not match:
        # même producteur, même millésime, pas de cuvée nommée de part et d'autre
        # (orthographes différentes de l'appellation : « Robola of Cephalonia » / « de Céphalonie »)
        for c in bouteilles:
            if (slug(c["producteur"]) == slug(v["producteur"])
                    and c["millesime"] == v["millesime"]
                    and not slug(c["cuvee"]) and not slug(v["cuvee"])
                    and c["couleur"] == v["couleur"]
                    and slug(c["appellation"]).split()[:1] == slug(v["appellation"]).split()[:1]):
                match = c
                break
    if match:
        match["prix_verre"] = v["prix_verre"]
        match["au_verre"] = True
        if v["section"].endswith("Apéritif"):
            match["aperitif"] = True
    else:
        v["au_verre"] = True
        final.append(v)

for v in final:
    v.setdefault("au_verre", False)
    v.setdefault("aperitif", False)

# ordre : on garde l'ordre de la carte
final.sort(key=lambda v: v["ligne"])
for n, v in enumerate(final, 1):
    v["id"] = n
    # champ de recherche agrégé
    v["q"] = slug(" ".join([
        v["label"], v["producteur"], v["cuvee"], v["appellation"], v["millesime"],
        v["origine"], v["region"], v["couleur"], v["cepages"], v["notes"],
        v["agence"], v["code"], v["format"], v.get("pays",""),
        "au verre" if v["au_verre"] else "",
    ]))
    del v["ligne"]

PAYS = ["France", "Italie", "Espagne", "Portugal", "Grèce", "Allemagne", "Autriche",
        "Hongrie", "République Tchèque", "Serbie", "Canada", "Québec", "États-Unis",
        "Californie", "Orégon", "Oregon", "Afrique du Sud", "Australie", "Nouvelle-Zélande",
        "Chili", "Argentine", "Liban", "Slovénie", "Suisse", "Géorgie"]

FR_REGIONS = ["Loire", "Bourgogne", "Beaujolais", "Bordeaux", "Rhône", "Alsace", "Jura",
              "Savoie", "Bugey", "Champagne", "Languedoc", "Roussillon", "Provence",
              "Corse", "Sud-Ouest", "Lang."]

CA_HINTS = ["Québec", "Quebec", "Ontario", "Niagara", "Dunham", "Brigham", "Farnham",
            "Saint-Armand", "Beamsville", "Twenty Mile"]


def devine_pays(v):
    orig = v["origine"]
    for p in PAYS:
        if slug(p) and slug(p) in slug(orig):
            return {"Québec": "Canada", "Californie": "États-Unis", "Orégon": "États-Unis",
                    "Oregon": "États-Unis"}.get(p, p)
    for h in CA_HINTS:
        if slug(h) in slug(orig) or slug(h) in slug(v["label"]):
            return "Canada"
    # le libellé du vin prime sur le titre de section, qui peut citer plusieurs pays
    proche = slug(v["origine"] + " " + v["appellation"])
    APPELLATION_PAYS = {
        "wachau": "Autriche", "kamptal": "Autriche", "kremstal": "Autriche",
        "weinviertel": "Autriche", "burgenland": "Autriche", "eisenberg": "Autriche",
        "weinland": "Autriche", "mosel": "Allemagne", "pfalz": "Allemagne",
        "pfalzer": "Allemagne", "rheinhessen": "Allemagne", "rheingau": "Allemagne",
        "nahe": "Allemagne", "ahr": "Allemagne", "baden": "Allemagne",
        "oberrhein": "Allemagne", "tokaji": "Hongrie", "moravie": "République Tchèque",
    }
    for k, pays_k in APPELLATION_PAYS.items():
        if k in proche.split() or k in proche:
            return pays_k
    for p in PAYS + ["Slovaquie"]:
        if slug(p) in proche and p != "France":
            return {"Québec": "Canada", "Californie": "États-Unis", "Orégon": "États-Unis",
                    "Oregon": "États-Unis"}.get(p, p)
    hay = slug(v["origine"] + " " + v["region"] + " " + v["appellation"])
    for p in PAYS + ["Slovaquie"]:
        if slug(p) in hay and p != "France":
            return {"Québec": "Canada", "Californie": "États-Unis", "Orégon": "États-Unis",
                    "Oregon": "États-Unis"}.get(p, p)
    toks = hay.split()
    for it in ["franciacorta", "etna", "sicile", "piemont", "toscane"]:
        if it in hay:
            return "Italie"
    if "it" in toks:          # « (It.) » abrégé dans le libellé
        return "Italie"
    for fr in FR_REGIONS + ["Touraine", "Anjou", "Cahors", "Vouvray", "Saumur", "Chablis",
                            "Beaune", "Nuits", "Mâcon", "Sancerre", "Muscadet", "Vin de France",
                            "Côtes du", "Vin de Pays", "Pernand", "Meursault", "Puligny", "Chassagne", "Gevrey", "Volnay", "Pommard", "Morey", "Vosne", "Chambolle", "Marsannay", "Mercurey", "Rully", "Bouzeron", "Fixin", "Aloxe", "Savigny", "Santenay", "Saint-Aubin", "Auxey", "Monthelie", "Ladoix", "Corton"]:
        if slug(fr) in hay:
            return "France"
    return ""


for v in final:
    v["region"] = (v["region"].replace("Bourgogne blanc", "Bourgogne")
                   .replace("Lang.-Roussillon", "Languedoc-Roussillon")
                   .replace("Péninsule Ibérique blanc", "Péninsule Ibérique"))
    v["pays"] = devine_pays(v)

meta = dict(
    titre="Carte des vins",
    total=len(final),
    au_verre=sum(1 for v in final if v["au_verre"]),
    couleurs=sorted({v["couleur"] for v in final if v["couleur"]}),
    regions=sorted({v["region"] for v in final if v["region"]}),
)

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(dict(meta=meta, vins=final), f, ensure_ascii=False, indent=1)

# data.js : ce que l'app charge réellement (champs de recherche recalculés côté navigateur)
compact = []
for v in final:
    w = {k: x for k, x in v.items() if k not in ("q", "section")}
    if w.get("notes") == w.get("cepages"):
        w["notes"] = ""
    compact.append(w)
with open("data.js", "w", encoding="utf-8") as f:
    f.write("window.CARTE = ")
    json.dump(dict(meta=meta, vins=compact), f, ensure_ascii=False, separators=(",", ":"))
    f.write(";\n")

print(json.dumps(meta, ensure_ascii=False, indent=1))
